"""scripts/build_slides.py: the deck.md parser, the 72-slide validation, the open-items report,
diagram caching, the PPTX writer, and (marked `browser`) a 2-slide PDF render."""

from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location("build_slides", REPO / "scripts" / "build_slides.py")
assert _spec and _spec.loader
bs = importlib.util.module_from_spec(_spec)
sys.modules["build_slides"] = bs
_spec.loader.exec_module(bs)

PENDING_LINE = "[[PENDING: S1 dots per arm | source: M8 batch 1 | how: make report]]"

FIXTURE = f"""<!--
Header comment: ignored by the parser, even with lines like
## 99. Not a slide
-->

::: diagram flow
flowchart LR
  A["Ticket"] --> B["Agent"]
  B --> C["Referee"]
:::

## 1. First slide
section: Opening
minutes: 0.5
layout: title
subtitle: A subtitle line

- One bullet
- A bullet that goes on
  over a second line

notes:
Say hello.

Second paragraph of notes.

## 2. Results
section: S1 Instructions
minutes: 7
diagram: flow now=B
slot: s1_dots
slot: s1_swap = 3 of 3
live: tonight's two dots

- The A/B verdicts
- Swap test verdict

{PENDING_LINE}

notes:
Read the dots.
[[PENDING: wall time of make demo-1 | source: M6 pilot | how: time make demo-1]]
"""


def parse(text: str = FIXTURE):
    return bs.parse_deck(text)


# -- parser -------------------------------------------------------------------------------------


def test_parses_blocks_with_numbers_titles_sections_and_minutes() -> None:
    deck = parse()
    assert [s.number for s in deck.slides] == [1, 2]
    first, second = deck.slides
    assert first.title == "First slide"
    assert first.section == "Opening"
    assert first.minutes == 0.5
    assert first.layout == "title"
    assert first.subtitle == "A subtitle line"
    assert second.section == "S1 Instructions"
    assert second.minutes == 7
    assert second.layout == "default"


def test_header_comment_is_ignored() -> None:
    assert 99 not in [s.number for s in parse().slides]


def test_bullets_join_continuation_lines() -> None:
    assert parse().slides[0].bullets == ["One bullet", "A bullet that goes on over a second line"]


def test_notes_keep_paragraphs_and_end_at_the_next_slide() -> None:
    notes = parse().slides[0].notes
    assert notes == "Say hello.\n\nSecond paragraph of notes."


def test_result_slots_are_named_and_empty_until_filled() -> None:
    slots = parse().slides[1].slots
    assert [(s.name, s.value, s.empty) for s in slots] == [
        ("s1_dots", "", True),
        ("s1_swap", "3 of 3", False),
    ]


def test_live_items_are_kept() -> None:
    assert parse().slides[1].live == ["tonight's two dots"]


def test_pending_markers_in_body_and_notes_are_detected() -> None:
    slide = parse().slides[1]
    assert [p.what for p in slide.pending] == ["S1 dots per arm"]
    assert slide.pending[0].source == "M8 batch 1"
    assert slide.pending[0].how == "make report"
    assert [p.what for p in slide.notes_pending] == ["wall time of make demo-1"]
    assert "PENDING" not in " ".join(slide.bullets)


def test_named_diagram_is_resolved_with_the_highlight() -> None:
    [source] = parse().slides[1].diagrams
    assert 'B["Agent (now)"]' in source
    assert 'A["Ticket"]' in source
    assert "class B now" in source
    assert "classDef now" in source


def test_inline_mermaid_and_code_and_image_fields() -> None:
    text = """## 1. Code
section: S3 Scope and sandbox
code: S3:policy 7-16
code: S4:filters 20-23, 30-36
image: images/board.png

```mermaid
flowchart LR
  X --> Y
```

- One
"""
    slide = parse(text).slides[0]
    assert slide.diagrams == ["flowchart LR\n  X --> Y"]
    assert slide.images == ["images/board.png"]
    assert [(c.anchor, c.ranges) for c in slide.code] == [
        ("S3:policy", [(7, 16)]),
        ("S4:filters", [(20, 23), (30, 36)]),
    ]


def test_unknown_text_in_a_slide_is_a_parse_error_with_its_line() -> None:
    text = "## 1. A\nsection: Opening\n\nstray paragraph\n"
    with pytest.raises(bs.ParseError, match="line 4"):
        parse(text)


def test_unknown_diagram_name_is_a_parse_error() -> None:
    with pytest.raises(bs.ParseError, match="nosuch"):
        parse("## 1. A\nsection: Opening\ndiagram: nosuch now=X\n")


# -- validation ---------------------------------------------------------------------------------


def _deck_text(numbers: list[int], extra: str = "") -> str:
    blocks = [
        f"## {n}. Slide {n}\nsection: Opening\nminutes: 1\n{extra}\n- a\n- b\n- c\n\nnotes:\nN{n}\n"
        for n in numbers
    ]
    return "\n".join(blocks)


def test_valid_72_slide_deck_has_no_errors(tmp_path: Path) -> None:
    deck = parse(_deck_text(list(range(1, 73))))
    assert bs.validate(deck, tmp_path) == []


def test_wrong_count_is_an_error(tmp_path: Path) -> None:
    errors = bs.validate(parse(_deck_text(list(range(1, 72)))), tmp_path)
    assert any("72" in e and "71" in e for e in errors)


def test_numbers_out_of_order_or_duplicated_are_errors(tmp_path: Path) -> None:
    numbers = list(range(1, 73))
    numbers[4], numbers[5] = numbers[5], numbers[4]
    assert any("order" in e for e in bs.validate(parse(_deck_text(numbers)), tmp_path))
    numbers = list(range(1, 72)) + [71]
    assert any("order" in e for e in bs.validate(parse(_deck_text(numbers)), tmp_path))


def test_missing_image_is_an_error_and_an_existing_one_is_not(tmp_path: Path) -> None:
    (tmp_path / "images").mkdir()
    (tmp_path / "images" / "ok.png").write_bytes(b"png")
    deck = parse(_deck_text([1, 2], extra="image: images/ok.png"))
    assert bs.validate(deck, tmp_path, expected=2) == []
    deck = parse(_deck_text([1, 2], extra="image: images/missing.png"))
    errors = bs.validate(deck, tmp_path, expected=2)
    assert len(errors) == 2 and all("images/missing.png" in e for e in errors)


def test_filled_slot_with_an_image_path_must_exist(tmp_path: Path) -> None:
    deck = parse(_deck_text([1], extra="slot: shot = images/shot.png"))
    assert any("images/shot.png" in e for e in bs.validate(deck, tmp_path, expected=1))


def test_more_than_six_bullets_or_twelve_code_lines_are_errors(tmp_path: Path) -> None:
    text = "## 1. Many\nsection: Opening\n" + "".join(f"- b{i}\n" for i in range(7))
    assert any("bullets" in e for e in bs.validate(parse(text), tmp_path, expected=1))
    text = "## 1. Code\nsection: Opening\ncode: S3:policy 1-13\n- a\n"
    assert any("code lines" in e for e in bs.validate(parse(text), tmp_path, expected=1))


def test_missing_section_and_unknown_anchor_are_errors(tmp_path: Path) -> None:
    deck = parse("## 1. A\n- a\n")
    assert any("section" in e for e in bs.validate(deck, tmp_path, expected=1))
    deck = parse("## 1. A\nsection: Opening\ncode: S9:nosuch 1-2\n")
    assert any("S9:nosuch" in e for e in bs.validate(deck, tmp_path, expected=1))


def test_open_items_list_every_pending_marker_and_empty_slot() -> None:
    pending, empty = bs.open_items(parse())
    assert [(n, where, p.what) for n, where, p in pending] == [
        (2, "slide", "S1 dots per arm"),
        (2, "notes", "wall time of make demo-1"),
    ]
    assert empty == [(2, "s1_dots")]


def test_check_mode_prints_counts_and_strict_fails_on_open_items(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    deck_md = tmp_path / "deck.md"
    deck_md.write_text(_deck_text([1, 2], extra=f"slot: x\n\n{PENDING_LINE}\n"), encoding="utf-8")
    argv = ["--deck", str(deck_md), "--expect", "2", "--check"]
    assert bs.main(argv) == 0
    out = capsys.readouterr().out
    assert "2 PENDING markers" in out and "2 empty result slots" in out
    assert bs.main([*argv, "--strict"]) == 1
    deck_md.write_text(_deck_text([1, 2]), encoding="utf-8")
    assert bs.main([*argv, "--strict"]) == 0


def test_check_mode_exits_2_on_validation_errors(tmp_path: Path) -> None:
    deck_md = tmp_path / "deck.md"
    deck_md.write_text(_deck_text([1, 3]), encoding="utf-8")
    assert bs.main(["--deck", str(deck_md), "--expect", "2", "--check"]) == 2


# -- code excerpts and diagrams -----------------------------------------------------------------


def test_code_excerpt_selects_lines_and_folds_skipped_content() -> None:
    ref = bs.CodeRef("S4:filters", [(20, 23), (30, 36)])
    excerpt = bs.resolve_code(ref, REPO)
    assert excerpt.path == "src/shipharness/filters/pipeline.py"
    assert excerpt.lines[0].strip() == "for flt in filters:"
    assert bs.FOLD in excerpt.lines
    assert excerpt.count == 11
    assert excerpt.text.splitlines()[0] == "for flt in filters:"  # common indent removed


def test_diagram_png_is_rendered_once_and_cached_by_content(tmp_path: Path) -> None:
    calls: list[str] = []

    def fake_render(source: str, out: Path) -> None:
        calls.append(source)
        out.write_bytes(b"\x89PNG fake")

    first = bs.diagram_png("flowchart LR\n A-->B", tmp_path, fake_render)
    again = bs.diagram_png("flowchart LR\n A-->B", tmp_path, fake_render)
    other = bs.diagram_png("flowchart LR\n A-->C", tmp_path, fake_render)
    assert first == again != other
    assert len(calls) == 2
    assert first.parent == tmp_path and first.suffix == ".png"


# -- the real deck ------------------------------------------------------------------------------


def test_the_course_deck_is_valid_and_matches_the_240_minute_schedule() -> None:
    deck_md = REPO / "slides" / "deck.md"
    deck = bs.parse_deck(deck_md.read_text(encoding="utf-8"))
    assert bs.validate(deck, deck_md.parent) == []
    assert sum(s.minutes or 0 for s in deck.slides) == 240
    assert all(s.notes for s in deck.slides)


# -- PPTX ---------------------------------------------------------------------------------------


def _png(path: Path, size: tuple[int, int] = (400, 200)) -> Path:
    from PIL import Image

    Image.new("RGB", size, "white").save(path)
    return path


def test_pptx_has_one_slide_per_block_with_titles_notes_and_pending(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = parse()
    pngs = {src: _png(tmp_path / f"d{i}.png") for i, src in enumerate(deck.slides[1].diagrams)}
    out = tmp_path / "deck.pptx"
    bs.write_pptx(bs.prepare(deck, tmp_path, REPO, pngs), out)
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    texts = [
        " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides
    ]
    assert "First slide" in texts[0] and "Results" in texts[1]
    assert "PENDING: S1 dots per arm" in texts[1]
    assert "s1_dots" in texts[1]
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
    assert "Say hello." in notes[0]
    assert "Read the dots." in notes[1] and "PENDING: wall time" in notes[1]
    assert "7 min" in notes[1]
    assert any(sh.shape_type == 13 for sh in prs.slides[1].shapes)  # the diagram picture


def test_pptx_code_slide_uses_a_monospace_box(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = parse("## 1. Policy\nsection: S3 Scope and sandbox\ncode: S3:policy 7-16\n- a\n")
    out = tmp_path / "code.pptx"
    bs.write_pptx(bs.prepare(deck, tmp_path, REPO, {}), out)
    shapes = Presentation(str(out)).slides[0].shapes
    code = [
        sh for sh in shapes if sh.has_text_frame and 'network: str = "none"' in sh.text_frame.text
    ]
    assert code
    run = code[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == bs.PPTX_MONO


# -- HTML and PDF -------------------------------------------------------------------------------


def test_html_has_one_page_per_slide_header_footer_and_pending_box(tmp_path: Path) -> None:
    html = bs.render_html(bs.prepare(parse(), tmp_path, REPO, {}))
    assert html.count('class="slide') == 2
    assert bs.COURSE in html
    assert "2 / 2" in html
    assert re.search(r'class="pending".*PENDING: S1 dots per arm', html, re.S)
    assert "s1_dots" in html


@pytest.mark.browser
def test_two_slide_fixture_renders_to_a_two_page_pdf(tmp_path: Path) -> None:
    pytest.importorskip("playwright.sync_api")
    text = FIXTURE.replace("diagram: flow now=B\n", "")
    slides = bs.prepare(parse(text), tmp_path, REPO, {})
    html_path = tmp_path / "deck.html"
    html_path.write_text(bs.render_html(slides), encoding="utf-8")
    pdf = tmp_path / "deck.pdf"
    try:
        report = bs.export_pdf(html_path, pdf)
    except bs.BrowserUnavailable as exc:
        pytest.skip(str(exc))
    assert bs.pdf_page_count(pdf) == 2
    assert set(report) == {1, 2}
