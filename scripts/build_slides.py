"""Build the course slides from slides/deck.md: an HTML deck, a PDF, and a PPTX with speaker notes.

Usage:
  uv run python scripts/build_slides.py            validate, then write slides/dist/deck.*
  uv run python scripts/build_slides.py --check    validate and list open items (no browser)
  uv run python scripts/build_slides.py --strict   exit 1 while a PENDING or CAPTURE marker or an
                                                   empty slot remains

The deck format is described in slides/README.md and in the comment at the top of slides/deck.md.
Mermaid diagrams are rendered to PNG once with Playwright Chromium (mermaid loaded from a pinned CDN
URL) and cached under slides/dist/diagrams by content hash. The HTML page is the layout engine: the
PDF is printed from it, and the PPTX uses the font sizes the HTML needed to fit each slide.

Exit codes: 0 built (open items are listed but allowed), 1 open items remain under --strict,
2 the deck does not parse or does not validate, 3 the browser step failed.
"""

from __future__ import annotations

import argparse
import functools
import hashlib
import html
import importlib.util
import re
import sys
import textwrap
import time
from collections.abc import Callable, Iterator
from contextlib import ExitStack, contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

REPO = Path(__file__).resolve().parents[1]
DECK = REPO / "slides" / "deck.md"
OUT = REPO / "slides" / "dist"
COURSE = "Module 1: Harness Engineering for Coding Agents"
EXPECTED = 72
MAX_BULLETS = 6
MAX_CODE_LINES = 12
MERMAID_URL = "https://cdn.jsdelivr.net/npm/mermaid@11.4.1/dist/mermaid.min.js"
# Bump when the diagram theme changes, so cached PNGs are rendered again.
DIAGRAM_THEME = "v2"
LAYOUTS = ("default", "title", "statement")
IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".svg", ".gif", ".webp")
FOLD = None  # marks left-out lines in a code excerpt
FOLD_TEXT = "..."
PPTX_FONT = "Calibri"
PPTX_MONO = "Consolas"
WIDTH, HEIGHT = 1280, 720  # px; 16:9, 13.333 x 7.5 in at 96 px per inch

_SLIDE = re.compile(r"^## (\d+)\.\s+(\S.*?)\s*$")
_FIELD = re.compile(r"^(section|minutes|layout|subtitle|diagram|image|code|slot|live):\s*(.*?)\s*$")
# [[PENDING: ...]] marks an open item; [[CAPTURE: ...]] marks output still to be copied in from a
# real run. Both are open items: listed in the report and failing --strict.
_PENDING = re.compile(r"^\s*\[\[(PENDING|CAPTURE):\s*(.+?)\s*\]\]\s*$")
_DIAGRAM_DEF = re.compile(r"^::: diagram ([A-Za-z0-9_-]+)\s*$")
_SLOT = re.compile(r"^([A-Za-z0-9_-]+)\s*(?:=\s*(.*))?$")
_RANGE = re.compile(r"^(\d+)(?:-(\d+))?$")


class ParseError(ValueError):
    """deck.md does not follow the block format; the message names the line."""


class BrowserUnavailable(RuntimeError):
    """Playwright or its Chromium is missing (uv run playwright install chromium)."""


@dataclass
class Pending:
    what: str
    source: str
    how: str
    raw: str
    kind: str = "PENDING"  # or "CAPTURE"


@dataclass
class Slot:
    name: str
    value: str

    @property
    def empty(self) -> bool:
        return not self.value


@dataclass
class CodeRef:
    anchor: str
    ranges: list[tuple[int, int]]


@dataclass
class Slide:
    number: int
    title: str
    line: int
    section: str = ""
    minutes: float | None = None
    layout: str = "default"
    subtitle: str = ""
    bullets: list[str] = field(default_factory=list)
    diagrams: list[str] = field(default_factory=list)
    images: list[str] = field(default_factory=list)
    code: list[CodeRef] = field(default_factory=list)
    slots: list[Slot] = field(default_factory=list)
    live: list[str] = field(default_factory=list)
    pending: list[Pending] = field(default_factory=list)
    notes: str = ""
    notes_pending: list[Pending] = field(default_factory=list)


@dataclass
class Deck:
    slides: list[Slide]
    diagrams: dict[str, str]


# -- parsing ------------------------------------------------------------------------------------


def _blank_comments(text: str) -> str:
    """Replace HTML comments with blank lines, keeping line numbers for error messages."""
    return re.sub(r"<!--.*?-->", lambda m: "\n" * m.group(0).count("\n"), text, flags=re.S)


def parse_pending(raw: str) -> Pending | None:
    match = _PENDING.match(raw)
    if not match:
        return None
    parts = [p.strip() for p in match.group(2).split(" | ")]
    info = {"source": "", "how": ""}
    for part in parts[1:]:
        key, _, value = part.partition(":")
        if key.strip() in info:
            info[key.strip()] = value.strip()
    return Pending(
        what=parts[0], source=info["source"], how=info["how"], raw=raw.strip(), kind=match.group(1)
    )


def _parse_ranges(spec: str, number: int) -> list[tuple[int, int]]:
    ranges = []
    for part in spec.replace(" ", "").split(","):
        match = _RANGE.match(part)
        if not match:
            raise ParseError(f"line {number}: code lines must look like 7-16 or 1-4,9-12")
        start = int(match.group(1))
        ranges.append((start, int(match.group(2) or start)))
    return ranges


def highlight(source: str, nodes: list[str], number: int) -> str:
    """The diagram with `nodes` marked: a thicker border (class `now`) and "(now)" in the label."""
    out = source
    for node in nodes:
        pattern = re.compile(rf'(\b{re.escape(node)}\[")([^"]*)("\])')
        if not pattern.search(out):
            raise ParseError(f"line {number}: node {node} is not defined in the diagram")
        out = pattern.sub(lambda m: f"{m.group(1)}{m.group(2)} (now){m.group(3)}", out, count=1)
    if "classDef now" not in out:
        out += "\n  classDef now stroke-width:5px"
    return out + f"\n  class {','.join(nodes)} now"


def parse_deck(text: str) -> Deck:
    """Parse deck.md into slides. Raises ParseError with the line number of the first problem."""
    lines = _blank_comments(text.replace("\r\n", "\n")).split("\n")
    slides: list[Slide] = []
    diagrams: dict[str, str] = {}
    refs: list[tuple[Slide, str, list[str], int]] = []
    slide: Slide | None = None
    in_notes = False
    notes: list[str] = []
    last_was_bullet = False
    i = 0

    def finish() -> None:
        if slide is not None:
            slide.notes = re.sub(r"\n{3,}", "\n\n", "\n".join(notes)).strip()

    while i < len(lines):
        line, number = lines[i], i + 1
        i += 1
        heading = _SLIDE.match(line)
        if heading:
            finish()
            slide = Slide(number=int(heading.group(1)), title=heading.group(2), line=number)
            slides.append(slide)
            in_notes, notes, last_was_bullet = False, [], False
            continue
        definition = _DIAGRAM_DEF.match(line)
        if definition and not in_notes:
            body = []
            while i < len(lines) and lines[i].strip() != ":::":
                body.append(lines[i])
                i += 1
            if i >= len(lines):
                raise ParseError(f"line {number}: diagram {definition.group(1)} has no closing :::")
            i += 1
            diagrams[definition.group(1)] = "\n".join(body).strip()
            continue
        if slide is None:
            if line.strip():
                raise ParseError(f"line {number}: text before the first slide ('## 1. Title')")
            continue
        if in_notes:
            notes.append(line.rstrip())
            pending = parse_pending(line)
            if pending:
                slide.notes_pending.append(pending)
            continue
        if not line.strip():
            last_was_bullet = False
            continue
        if line.strip() == "notes:":
            in_notes = True
            continue
        if line.startswith("- "):
            slide.bullets.append(line[2:].strip())
            last_was_bullet = True
            continue
        if line[:1].isspace() and last_was_bullet:
            slide.bullets[-1] += " " + line.strip()
            continue
        last_was_bullet = False
        pending = parse_pending(line)
        if pending:
            slide.pending.append(pending)
            continue
        if line.strip() == "```mermaid":
            body = []
            while i < len(lines) and lines[i].strip() != "```":
                body.append(lines[i])
                i += 1
            if i >= len(lines):
                raise ParseError(f"line {number}: mermaid block has no closing ```")
            i += 1
            slide.diagrams.append("\n".join(body).strip())
            continue
        match = _FIELD.match(line)
        if not match:
            raise ParseError(
                f"line {number}: unexpected text in slide {slide.number}: {line.strip()[:60]!r} "
                "(bullets start with '- ', fields with 'name:', notes after 'notes:')"
            )
        key, value = match.groups()
        if key == "section":
            slide.section = value
        elif key == "minutes":
            try:
                slide.minutes = float(value)
            except ValueError as exc:
                raise ParseError(f"line {number}: minutes must be a number, got {value!r}") from exc
        elif key == "layout":
            if value not in LAYOUTS:
                raise ParseError(f"line {number}: layout must be one of {', '.join(LAYOUTS)}")
            slide.layout = value
        elif key == "subtitle":
            slide.subtitle = value
        elif key == "image":
            slide.images.append(value)
        elif key == "live":
            slide.live.append(value)
        elif key == "slot":
            slot = _SLOT.match(value)
            if not slot:
                raise ParseError(
                    f"line {number}: slot must look like 'slot: name' or 'name = value'"
                )
            slide.slots.append(Slot(slot.group(1), (slot.group(2) or "").strip()))
        elif key == "code":
            anchor, _, spec = value.partition(" ")
            if not spec.strip():
                raise ParseError(
                    f"line {number}: code needs an anchor id and lines, e.g. S3:policy 7-16"
                )
            slide.code.append(CodeRef(anchor, _parse_ranges(spec, number)))
        elif key == "diagram":
            name, _, rest = value.partition(" ")
            nodes: list[str] = []
            if rest.strip():
                if not rest.strip().startswith("now="):
                    raise ParseError(f"line {number}: diagram highlight must look like now=A,B")
                nodes = [n for n in rest.strip()[4:].split(",") if n]
            refs.append((slide, name, nodes, number))
            slide.diagrams.append("")  # filled below, once every definition is known
    finish()
    for owner, name, nodes, number in refs:
        if name not in diagrams:
            raise ParseError(
                f"line {number}: no diagram named {name!r} (define it with ::: diagram)"
            )
        resolved = highlight(diagrams[name], nodes, number) if nodes else diagrams[name]
        owner.diagrams[owner.diagrams.index("")] = resolved
    return Deck(slides=slides, diagrams=diagrams)


# -- code excerpts ------------------------------------------------------------------------------


@dataclass
class CodeExcerpt:
    anchor: str
    path: str
    lines: list[str | None]  # raw lines; FOLD (None) where content was left out
    count: int  # selected lines, folds not counted

    @property
    def text(self) -> str:
        """The excerpt with its common indent removed; a fold shows as '...' one level in."""
        real = [line for line in self.lines if line is not FOLD]
        body = iter(textwrap.dedent("\n".join(real)).split("\n"))
        out: list[str] = []
        prev = ""
        for line in self.lines:
            if line is FOLD:
                indent = len(prev) - len(prev.lstrip()) + (4 if prev.rstrip().endswith(":") else 0)
                out.append(" " * indent + FOLD_TEXT)
            else:
                prev = next(body)
                out.append(prev)
        return "\n".join(out)


@functools.lru_cache(maxsize=4)
def _anchor_index(root: Path) -> dict[str, tuple[str, list[str]]]:
    spec = importlib.util.spec_from_file_location("_anchors", REPO / "scripts" / "anchors.py")
    assert spec and spec.loader
    anchors = importlib.util.module_from_spec(spec)
    sys.modules["_anchors"] = anchors  # dataclasses look the module up while it loads
    spec.loader.exec_module(anchors)
    index = {}
    for anchor in anchors.collect(root)[0]:
        lines = anchor.path.read_text(encoding="utf-8").splitlines()
        index[anchor.anchor_id] = (
            anchor.path.relative_to(root).as_posix(),
            lines[anchor.start : anchor.end - 1],
        )
    return index


def resolve_code(ref: CodeRef, root: Path = REPO) -> CodeExcerpt:
    """The selected lines of a walkthrough anchor. Raises KeyError or ValueError when invalid."""
    index = _anchor_index(Path(root))
    if ref.anchor not in index:
        raise KeyError(f"no walkthrough anchor named {ref.anchor}")
    path, source = index[ref.anchor]
    lines: list[str | None] = []
    count = 0
    previous_end = None
    for start, end in ref.ranges:
        if not 1 <= start <= end <= len(source):
            raise ValueError(f"{ref.anchor} lines {start}-{end} outside 1-{len(source)}")
        if previous_end is not None and any(s.strip() for s in source[previous_end : start - 1]):
            lines.append(FOLD)
        lines.extend(source[start - 1 : end])
        count += end - start + 1
        previous_end = end
    return CodeExcerpt(anchor=ref.anchor, path=path, lines=lines, count=count)


# -- validation and open items ------------------------------------------------------------------


def _is_image(value: str) -> bool:
    return value.lower().endswith(IMAGE_SUFFIXES)


def validate(deck: Deck, base_dir: Path, expected: int = EXPECTED, repo: Path = REPO) -> list[str]:
    """Every problem that stops a build: count, numbering, missing images, limits, bad code refs."""
    errors: list[str] = []
    numbers = [s.number for s in deck.slides]
    if len(numbers) != expected:
        errors.append(f"the deck has {len(numbers)} slides; it must have exactly {expected}")
    for position, number in enumerate(numbers, start=1):
        if number != position:
            errors.append(
                f"slide numbers must run 1-{len(numbers)} in order: position {position} is "
                f"numbered {number} (line {deck.slides[position - 1].line})"
            )
            break
    for s in deck.slides:
        where = f"slide {s.number} (line {s.line})"
        if not s.section:
            errors.append(f"{where}: no section")
        if len(s.bullets) > MAX_BULLETS:
            errors.append(f"{where}: {len(s.bullets)} bullets (max {MAX_BULLETS})")
        for image in s.images + [x.value for x in s.slots if _is_image(x.value)]:
            if not (base_dir / image).is_file():
                errors.append(f"{where}: image not found: {image} (paths are relative to the deck)")
        total = 0
        for ref in s.code:
            try:
                total += resolve_code(ref, repo).count
            except (KeyError, ValueError) as exc:
                errors.append(f"{where}: code {ref.anchor}: {exc.args[0]}")
        if total > MAX_CODE_LINES:
            errors.append(f"{where}: {total} code lines (max {MAX_CODE_LINES})")
    return errors


def open_items(deck: Deck) -> tuple[list[tuple[int, str, Pending]], list[tuple[int, str]]]:
    """Every PENDING or CAPTURE marker (on the slide or in the notes) and every empty slot."""
    pending: list[tuple[int, str, Pending]] = []
    empty: list[tuple[int, str]] = []
    for s in deck.slides:
        pending += [(s.number, "slide", p) for p in s.pending]
        pending += [(s.number, "notes", p) for p in s.notes_pending]
        empty += [(s.number, x.name) for x in s.slots if x.empty]
    return pending, empty


# -- diagrams -----------------------------------------------------------------------------------


def diagram_png(source: str, cache_dir: Path, render: Callable[[str, Path], None]) -> Path:
    """The PNG for one Mermaid source, rendered only when no cached file has the same hash."""
    key = hashlib.sha256(f"{MERMAID_URL}\n{DIAGRAM_THEME}\n{source}".encode()).hexdigest()[:16]
    path = Path(cache_dir) / f"{key}.png"
    if not path.is_file():
        path.parent.mkdir(parents=True, exist_ok=True)
        render(source, path)
    return path


MERMAID_INIT = {
    "startOnLoad": False,
    "theme": "base",
    "themeVariables": {
        "fontFamily": "Ubuntu Sans, DejaVu Sans, Liberation Sans, sans-serif",
        "fontSize": "22px",
        "primaryColor": "#eef3f8",
        "primaryBorderColor": "#1f3a5f",
        "primaryTextColor": "#0b1f33",
        "lineColor": "#1f3a5f",
        "secondaryColor": "#f6f8fa",
        "tertiaryColor": "#ffffff",
    },
    "flowchart": {
        "htmlLabels": True,
        "useMaxWidth": False,
        "padding": 8,
        "nodeSpacing": 16,
        "rankSpacing": 36,
        "wrappingWidth": 260,
    },
}


@contextmanager
def _browser() -> Iterator[Any]:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise BrowserUnavailable(f"playwright is not installed: {exc}") from exc
    manager = sync_playwright().start()
    try:
        try:
            browser = manager.chromium.launch()
        except Exception as exc:
            raise BrowserUnavailable(
                f"Chromium did not start (uv run playwright install chromium): {exc}"
            ) from exc
        try:
            yield browser
        finally:
            browser.close()
    finally:
        manager.stop()


def render_diagrams(sources: list[str], cache_dir: Path) -> dict[str, Path]:
    """Render every uncached Mermaid source to PNG in one browser session."""
    import json

    pngs: dict[str, Path] = {}
    pages: list[Any] = []

    with ExitStack() as stack:

        def render(source: str, out: Path) -> None:
            if not pages:  # the browser starts only when a diagram is not cached yet
                browser = stack.enter_context(_browser())
                page = browser.new_page(
                    device_scale_factor=2, viewport={"width": 2400, "height": 1600}
                )
                page.set_content("<!doctype html><html><body style='margin:0'></body></html>")
                page.add_script_tag(url=MERMAID_URL)
                page.evaluate(f"mermaid.initialize({json.dumps(MERMAID_INIT)})")
                pages.append(page)
            page = pages[0]
            page.evaluate(
                """async ([id, src]) => {
                    const {svg} = await mermaid.render(id, src);
                    document.body.innerHTML = '<div id="wrap" style="display:inline-block;'
                        + 'padding:8px;background:#fff">' + svg + '</div>';
                }""",
                [f"d{len(pngs)}", source],
            )
            page.locator("#wrap").screenshot(path=str(out))

        for source in dict.fromkeys(sources):
            pngs[source] = diagram_png(source, cache_dir, render)
    return pngs


# -- prepared slides (shared by HTML and PPTX) --------------------------------------------------


@dataclass
class Visual:
    path: Path | None  # None: a diagram that was not rendered
    size: tuple[int, int] | None
    label: str


@dataclass
class RenderSlide:
    slide: Slide
    total: int
    layout: str  # title, statement, text, split (text left, visual right), stacked
    code: list[CodeExcerpt]
    visuals: list[Visual]


def _image_size(path: Path) -> tuple[int, int] | None:
    try:
        from PIL import Image

        with Image.open(path) as img:
            return img.size
    except Exception:  # noqa: BLE001 - an unreadable image is shown without a known size
        return None


def prepare(
    deck: Deck, base_dir: Path, repo: Path = REPO, diagram_pngs: dict[str, Path] | None = None
) -> list[RenderSlide]:
    """Resolve code, diagrams, and images, and choose each slide's layout."""
    pngs = diagram_pngs or {}
    total = len(deck.slides)
    out = []
    for s in deck.slides:
        code = [resolve_code(ref, repo) for ref in s.code]
        visuals = []
        for source in s.diagrams:
            path = pngs.get(source)
            visuals.append(Visual(path, _image_size(path) if path else None, "diagram"))
        for image in s.images + [x.value for x in s.slots if _is_image(x.value)]:
            path = Path(base_dir) / image
            visuals.append(Visual(path, _image_size(path), image))
        if s.layout in ("title", "statement"):
            layout = s.layout
        elif code:
            layout = "stacked"
        elif visuals:
            size = visuals[0].size
            layout = "stacked" if size and size[0] / size[1] > 1.8 else "split"
        else:
            layout = "text"
        out.append(RenderSlide(s, total, layout, code, visuals))
    return out


# -- HTML ---------------------------------------------------------------------------------------

CSS = """
@page { size: 1280px 720px; margin: 0; }
html, body { margin: 0; padding: 0; background: #ffffff; }
body { font-family: "Ubuntu Sans", "DejaVu Sans", "Liberation Sans", Arial, sans-serif;
  color: #111820; }
.slide { width: 1280px; height: 720px; box-sizing: border-box; padding: 0 56px; position: relative;
  overflow: hidden; display: flex; flex-direction: column; background: #ffffff;
  break-after: page; page-break-after: always; }
.slide:last-child { break-after: auto; page-break-after: auto; }
header { flex: 0 0 52px; display: flex; justify-content: space-between; align-items: flex-end;
  padding-bottom: 8px; border-bottom: 3px solid #1f3a5f; font-size: 19px; color: #1f3a5f; }
header .section { font-weight: 700; letter-spacing: 0.02em; }
header .num { font-variant-numeric: tabular-nums; font-weight: 600; }
h1 { font-size: 40px; line-height: 1.15; margin: 20px 0 14px; color: #0b1f33; font-weight: 700; }
.subtitle { font-size: 28px; color: #2c3e50; margin: 0 0 18px; line-height: 1.3; }
.body { flex: 1 1 0; min-height: 0; display: flex; gap: 36px; font-size: 28px; overflow: hidden; }
.text { flex: 1 1 0; min-width: 0; min-height: 0; overflow: hidden; }
.visual { flex: 1.15 1 0; min-width: 0; min-height: 0; display: flex; flex-direction: column;
  justify-content: center; align-items: stretch; gap: 10px; overflow: hidden; }
.visual img { width: 100%; height: 100%; min-height: 0; object-fit: contain; }
.layout-stacked .body { flex-direction: column; gap: 14px; }
.layout-stacked .text { flex: 0 0 auto; }
.layout-stacked .visual { flex: 1 1 0; justify-content: flex-start; }
.layout-stacked .visual.has-img { min-height: 260px; }
ul { margin: 0; padding-left: 1.1em; }
li { margin: 0 0 0.5em; line-height: 1.28; overflow-wrap: anywhere; }
li::marker { color: #1f3a5f; }
code { font-family: "DejaVu Sans Mono", "Liberation Mono", monospace; font-size: 0.86em;
  background: #eef1f5; padding: 0 0.15em; border-radius: 3px; }
.codeblock { flex: 0 0 auto; }
pre { font-family: "DejaVu Sans Mono", "Liberation Mono", monospace; font-size: 21px;
  line-height: 1.32; margin: 0; padding: 12px 18px; background: #f4f6f8; color: #111820;
  border: 1px solid #c9d2dc; border-left: 6px solid #1f3a5f; white-space: pre; overflow: hidden; }
.codepath { font-size: 16px; color: #3d4b59; margin: 4px 0 0;
  font-family: "DejaVu Sans Mono", monospace; }
.placeholder { border: 2px dashed #8a99a8; color: #3d4b59; font-size: 20px; padding: 20px;
  text-align: center; }
.boxes { flex: 0 0 auto; display: flex; flex-wrap: wrap; align-items: flex-start; gap: 8px;
  margin: 10px 0 10px; }
.boxes > div { box-sizing: border-box; }
.pending { flex: 1 1 40%; background: #fff3cd; border: 3px solid #b26a00; color: #3d2600;
  border-radius: 6px; padding: 7px 14px; font-size: 20px; line-height: 1.25; }
.slot, .live { flex: 0 1 auto; }
.pending strong { font-weight: 700; }
.pending small { display: block; font-size: 15px; color: #5c4400; margin-top: 2px; }
.slot { border: 3px dashed #4d5d6e; color: #1f2a35; border-radius: 6px; padding: 6px 14px;
  font-size: 19px; }
.slot.filled { border-style: solid; border-color: #1f6f43; background: #eef8f1; font-size: 21px; }
.live { background: #eef1f5; border: 2px solid #7d8c9b; color: #1f2a35; border-radius: 6px;
  padding: 6px 14px; font-size: 19px; }
footer { flex: 0 0 38px; display: flex; justify-content: space-between; align-items: center;
  border-top: 1px solid #c9d2dc; font-size: 16px; color: #3d4b59; margin-top: auto; }
.layout-title .titlebox, .layout-statement .titlebox { flex: 1 1 0; display: flex;
  flex-direction: column; justify-content: center; align-items: center; text-align: center;
  min-height: 0; }
.layout-title h1 { font-size: 64px; margin: 0 0 24px; }
.layout-title .subtitle { font-size: 32px; max-width: 1000px; }
.layout-statement h1 { font-size: 60px; margin: 0 0 24px; }
.layout-title ul, .layout-statement ul { list-style: none; padding: 0; font-size: 28px;
  color: #2c3e50; }
"""

FIT_JS = """
() => {
  const report = {};
  const over = (el) => el.scrollHeight > el.clientHeight + 1 || el.scrollWidth > el.clientWidth + 1;
  for (const slide of document.querySelectorAll('section.slide')) {
    const n = Number(slide.dataset.n);
    const body = slide.querySelector('.body');
    if (!body) { report[n] = {font_px: 28, code_px: 0, overflow: false}; continue; }
    const pres = [...body.querySelectorAll('pre')];
    for (const pre of pres) {
      let c = parseFloat(getComputedStyle(pre).fontSize);
      while (pre.scrollWidth > pre.clientWidth + 1 && c > 13) {
        c -= 0.5;
        pre.style.fontSize = c + 'px';
      }
      pre.dataset.base = c;
    }
    const base = parseFloat(getComputedStyle(body).fontSize);
    const parts = () => [body, ...body.querySelectorAll('.text, .visual, pre')];
    let k = 1.0;
    while (parts().some(over) && k > 0.62) {
      k -= 0.03;
      body.style.fontSize = (base * k) + 'px';
      for (const pre of pres) pre.style.fontSize = (pre.dataset.base * Math.max(k, 0.75)) + 'px';
    }
    report[n] = {
      font_px: Math.round(base * k * 10) / 10,
      code_px: pres.length ? parseFloat(getComputedStyle(pres[0]).fontSize) : 0,
      overflow: parts().some(over),
    };
  }
  return report;
}
"""


def _inline(text: str) -> str:
    return re.sub(r"`([^`]+)`", r"<code>\1</code>", html.escape(text))


def _visual_html(v: Visual) -> str:
    if v.path is None:
        return '<div class="placeholder">diagram not rendered (run make slides)</div>'
    return f'<img src="{v.path.resolve().as_uri()}" alt="{html.escape(v.label)}">'


def _boxes_html(s: Slide) -> str:
    parts = []
    for p in s.pending:
        small = _pending_small(p)
        parts.append(
            f'<div class="pending"><strong>{p.kind}: {html.escape(p.what)}</strong>'
            + (f"<small>{html.escape(small)}</small>" if small else "")
            + "</div>"
        )
    for slot in s.slots:
        if slot.empty:
            parts.append(
                f'<div class="slot">Result slot <code>{html.escape(slot.name)}</code>: empty</div>'
            )
        elif not _is_image(slot.value):
            parts.append(
                f'<div class="slot filled"><code>{html.escape(slot.name)}</code>: '
                f"{_inline(slot.value)}</div>"
            )
    for item in s.live:
        parts.append(f'<div class="live">Filled in live: {_inline(item)}</div>')
    return f'<div class="boxes">{"".join(parts)}</div>' if parts else ""


def _slide_html(rs: RenderSlide) -> str:
    s = rs.slide
    head = (
        f'<header><span class="section">{html.escape(s.section)}</span>'
        f'<span class="num">{s.number} / {rs.total}</span></header>'
    )
    foot = f"<footer><span>{html.escape(COURSE)}</span><span></span></footer>"
    bullets = "".join(f"<li>{_inline(b)}</li>" for b in s.bullets)
    ul = f"<ul>{bullets}</ul>" if bullets else ""
    if rs.layout in ("title", "statement"):
        sub = f'<p class="subtitle">{_inline(s.subtitle)}</p>' if s.subtitle else ""
        middle = f'<div class="titlebox"><h1>{_inline(s.title)}</h1>{sub}{ul}</div>'
        return (
            f'<section class="slide layout-{rs.layout}" data-n="{s.number}">{head}{middle}'
            f"{_boxes_html(s)}{foot}</section>"
        )
    sub = f'<p class="subtitle">{_inline(s.subtitle)}</p>' if s.subtitle else ""
    visual_parts = [
        f'<div class="codeblock"><pre>{html.escape(c.text)}</pre>'
        f'<div class="codepath">{html.escape(c.path)}</div></div>'
        for c in rs.code
    ] + [_visual_html(v) for v in rs.visuals]
    visual = ""
    if visual_parts:
        cls = "visual has-img" if rs.visuals else "visual"
        visual = f'<div class="{cls}">{"".join(visual_parts)}</div>'
    text = f'<div class="text">{ul}</div>' if ul else ""
    return (
        f'<section class="slide layout-{rs.layout}" data-n="{s.number}">{head}'
        f'<h1>{_inline(s.title)}</h1>{sub}<div class="body">{text}{visual}</div>'
        f"{_boxes_html(s)}{foot}</section>"
    )


def render_html(slides: list[RenderSlide]) -> str:
    """One HTML page; each slide is a 1280x720 section that prints as one PDF page."""
    body = "\n".join(_slide_html(rs) for rs in slides)
    return (
        '<!doctype html>\n<html lang="en"><head><meta charset="utf-8">'
        f"<title>{html.escape(COURSE)}</title><style>{CSS}</style></head>"
        f"<body>\n{body}\n</body></html>\n"
    )


def export_pdf(html_path: Path, pdf_path: Path) -> dict[int, dict[str, Any]]:
    """Fit each slide's text (shrinking the font when needed), then print the PDF.
    Returns, per slide number, the body font size used and whether it still overflows."""
    with _browser() as browser:
        page = browser.new_page(viewport={"width": WIDTH, "height": HEIGHT})
        page.goto(Path(html_path).resolve().as_uri())
        page.wait_for_load_state("load")
        page.evaluate("document.fonts.ready")
        raw = page.evaluate(FIT_JS)
        page.pdf(
            path=str(pdf_path),
            width=f"{WIDTH}px",
            height=f"{HEIGHT}px",
            print_background=True,
            margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
            prefer_css_page_size=True,
        )
    return {int(k): v for k, v in raw.items()}


def pdf_page_count(path: Path) -> int:
    return len(re.findall(rb"/Type\s*/Page(?![s\w])", Path(path).read_bytes()))


# -- PPTX ---------------------------------------------------------------------------------------

NAVY, INK, MUTED = (0x1F, 0x3A, 0x5F), (0x11, 0x18, 0x20), (0x3D, 0x4B, 0x59)


def _px(value: float) -> int:
    return int(round(value * 9525))  # EMU per px at 96 dpi


def _add_text(
    slide: Any,
    box: tuple[float, float, float, float],
    paragraphs: list[str],
    size_pt: float,
    *,
    bold: bool = False,
    color: tuple[int, int, int] = INK,
    align: str = "left",
    mono: bool = False,
    bullets: bool = False,
    space_after_pt: float = 0,
) -> Any:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    x, y, w, h = box
    shape = slide.shapes.add_textbox(_px(x), _px(y), _px(w), _px(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for margin in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, margin, _px(2))
    for index, text in enumerate(paragraphs):
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.alignment = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }[align]
        if space_after_pt:
            para.space_after = Pt(space_after_pt)
        pieces = [text] if mono else re.split(r"(`[^`]+`)", text)
        if bullets:
            pieces = ["•  ", *pieces]
        for piece in pieces:
            if not piece:
                continue
            is_code = mono or (piece.startswith("`") and piece.endswith("`") and len(piece) > 1)
            run = para.add_run()
            run.text = piece.strip("`") if is_code and not mono else piece
            run.font.size = Pt(size_pt)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*color)
            run.font.name = PPTX_MONO if is_code else PPTX_FONT
    return shape


def _add_rect(
    slide: Any,
    box: tuple[float, float, float, float],
    fill: tuple[int, int, int] | None,
    line: tuple[int, int, int] | None = None,
    line_px: float = 0,
    dashed: bool = False,
) -> Any:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    x, y, w, h = box
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _px(x), _px(y), _px(w), _px(h))
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = RGBColor(*line)
        shape.line.width = Pt(line_px * 0.75)
        if dashed:
            shape.line.dash_style = MSO_LINE.DASH
    else:
        shape.line.fill.background()
    return shape


def _lines_needed(text: str, width_px: float, font_px: float, char_em: float = 0.5) -> int:
    per_line = max(1, int(width_px / (font_px * char_em)))
    return max(1, -(-len(text) // per_line))


def _add_picture(slide: Any, v: Visual, box: tuple[float, float, float, float]) -> None:
    x, y, w, h = box
    if v.path is None or not v.path.is_file():
        _add_text(slide, box, ["diagram not rendered (run make slides)"], 14, color=MUTED)
        return
    iw, ih = v.size or (w, h)
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(
        str(v.path), _px(x + (w - pw) / 2), _px(y + (h - ph) / 2), _px(pw), _px(ph)
    )


def _pending_small(p: Pending) -> str:
    parts = [f"source: {p.source}" if p.source else "", f"how: {p.how}" if p.how else ""]
    return " · ".join(x for x in parts if x)


def _box_layout(
    s: Slide, inner: float
) -> tuple[list[tuple[str, list[str], float, float, float, float]], float]:
    """The PENDING, slot, and live boxes as (kind, paragraphs, x, y, w, h), relative to the top
    of the box area, plus the area's height. PENDING boxes go two per row, like the HTML."""
    gap = 8.0
    placed: list[tuple[str, list[str], float, float, float, float]] = []
    y = 0.0
    pend = [
        ([f"{p.kind}: {p.what}"] + ([_pending_small(p)] if _pending_small(p) else []))
        for p in s.pending
    ]
    for i in range(0, len(pend), 2):
        row = pend[i : i + 2]
        w = (inner - gap * (len(row) - 1)) / len(row)
        h = max(
            16
            + _lines_needed(paras[0], w - 28, 20) * 26
            + (_lines_needed(paras[1], w - 28, 15) * 19 if len(paras) > 1 else 0)
            for paras in row
        )
        for j, paras in enumerate(row):
            placed.append(("pending", paras, j * (w + gap), y, w, h))
        y += h + gap
    chips: list[tuple[str, str]] = []
    for slot in s.slots:
        if slot.empty:
            chips.append(("slot", f"Result slot {slot.name}: empty"))
        elif not _is_image(slot.value):
            chips.append(("filled", f"{slot.name}: {slot.value}"))
    chips += [("live", f"Filled in live: {item}") for item in s.live]
    x, row_h = 0.0, 0.0
    for kind, text in chips:
        w = min(inner, len(text) * 19 * 0.52 + 36)
        if x and x + w > inner:
            x, y = 0.0, y + row_h + gap
            row_h = 0.0
        h = 14 + _lines_needed(text, w - 28, 19) * 26
        placed.append((kind, [text], x, y, w, h))
        x += w + gap
        row_h = max(row_h, h)
    height = y + row_h if chips else max(0.0, y - gap)
    return placed, height


def _pptx_slide(prs: Any, rs: RenderSlide, fit: dict[str, Any] | None) -> None:
    s = rs.slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pad, inner = 56, WIDTH - 112
    # Header: section left, number right, rule below.
    _add_text(slide, (pad, 14, inner * 0.7, 32), [s.section], 14, bold=True, color=NAVY)
    _add_text(
        slide,
        (pad + inner * 0.7, 14, inner * 0.3, 32),
        [f"{s.number} / {rs.total}"],
        14,
        bold=True,
        color=NAVY,
        align="right",
    )
    _add_rect(slide, (pad, 50, inner, 3), NAVY)
    # Footer.
    _add_rect(slide, (pad, HEIGHT - 38, inner, 1), (0xC9, 0xD2, 0xDC))
    _add_text(slide, (pad, HEIGHT - 32, inner, 26), [COURSE], 12, color=MUTED)

    body_px = float((fit or {}).get("font_px") or 28)
    code_px = float((fit or {}).get("code_px") or 21)
    boxes, boxes_h = _box_layout(s, inner)
    bottom = HEIGHT - 48 - (boxes_h + 10 if boxes else 0)

    if rs.layout in ("title", "statement"):
        title_px = 64 if rs.layout == "title" else 60
        lines = _lines_needed(s.title, inner, title_px, 0.52)
        paras_h = lines * title_px * 1.2
        sub_h = _lines_needed(s.subtitle, 1000, 32) * 42 if s.subtitle else 0
        list_h = len(s.bullets) * 40
        top = 60 + max(0, (bottom - 60 - paras_h - sub_h - list_h) / 2)
        _add_text(
            slide,
            (pad, top, inner, paras_h),
            [s.title],
            title_px * 0.75,
            bold=True,
            color=(0x0B, 0x1F, 0x33),
            align="center",
        )
        y = top + paras_h + 8
        if s.subtitle:
            _add_text(
                slide,
                (pad + 84, y, inner - 168, sub_h),
                [s.subtitle],
                24,
                color=(0x2C, 0x3E, 0x50),
                align="center",
            )
            y += sub_h + 8
        if s.bullets:
            _add_text(
                slide,
                (pad, y, inner, list_h),
                s.bullets,
                21,
                color=(0x2C, 0x3E, 0x50),
                align="center",
                space_after_pt=6,
            )
    else:
        title_lines = _lines_needed(s.title, inner, 40, 0.52)
        title_h = title_lines * 46 + 6
        _add_text(
            slide, (pad, 66, inner, title_h), [s.title], 30, bold=True, color=(0x0B, 0x1F, 0x33)
        )
        top = 66 + title_h + 10
        if s.subtitle:
            _add_text(slide, (pad, top, inner, 40), [s.subtitle], 21, color=(0x2C, 0x3E, 0x50))
            top += 46
        avail = max(60.0, bottom - top)
        visual_count = len(rs.code) + len(rs.visuals)
        if rs.layout == "split":
            text_w = (inner - 36) / 2.15
            text_box = (pad, top, text_w, avail)
            vis_box = (pad + text_w + 36, top, inner - text_w - 36, avail)
        else:
            text_w = inner
            line_h = body_px * 1.28
            text_h = sum(
                _lines_needed(b, inner - 30, body_px) * line_h + body_px * 0.5 for b in s.bullets
            )
            text_h = min(text_h, avail) if visual_count else avail
            text_box = (pad, top, inner, text_h)
            vis_box = (pad, top + text_h + 10, inner, max(40.0, avail - text_h - 10))
        if s.bullets:
            _add_text(
                slide,
                text_box,
                s.bullets,
                body_px * 0.75,
                bullets=True,
                space_after_pt=body_px * 0.75 * 0.45,
            )
        if visual_count:
            vx, vy, vw, vh = vis_box
            code_heights = [len(c.lines) * code_px * 1.32 + 24 for c in rs.code]
            pic_h = (vh - sum(code_heights) - 26 * len(rs.code)) / max(1, len(rs.visuals))
            for excerpt, ch in zip(rs.code, code_heights, strict=True):
                _add_rect(slide, (vx, vy, vw, ch), (0xF4, 0xF6, 0xF8), (0xC9, 0xD2, 0xDC), 1)
                _add_rect(slide, (vx, vy, 6, ch), NAVY)
                _add_text(
                    slide,
                    (vx + 18, vy + 10, vw - 28, ch - 16),
                    excerpt.text.split("\n"),
                    code_px * 0.75,
                    mono=True,
                )
                _add_text(
                    slide, (vx, vy + ch + 2, vw, 22), [excerpt.path], 11, color=MUTED, mono=True
                )
                vy += ch + 26
            for v in rs.visuals:
                _add_picture(slide, v, (vx, vy, vw, max(40.0, pic_h)))
                vy += pic_h

    top_of_boxes = HEIGHT - 48 - boxes_h
    styles = {
        "pending": ((0xFF, 0xF3, 0xCD), (0xB2, 0x6A, 0x00), False, (0x3D, 0x26, 0x00)),
        "slot": (None, (0x4D, 0x5D, 0x6E), True, (0x1F, 0x2A, 0x35)),
        "filled": ((0xEE, 0xF8, 0xF1), (0x1F, 0x6F, 0x43), False, (0x1F, 0x2A, 0x35)),
        "live": ((0xEE, 0xF1, 0xF5), (0x7D, 0x8C, 0x9B), False, (0x1F, 0x2A, 0x35)),
    }
    for kind, paras, bx, by, bw, bh in boxes:
        fill, line, dashed, color = styles[kind]
        box = (pad + bx, top_of_boxes + by, bw, bh)
        _add_rect(slide, box, fill, line, 3 if kind != "live" else 2, dashed)
        shape = _add_text(
            slide,
            (box[0] + 12, box[1] + 5, bw - 24, bh - 8),
            paras[:1],
            15 if kind == "pending" else 14,
            bold=kind == "pending",
            color=color,
        )
        if len(paras) > 1:
            from pptx.dml.color import RGBColor
            from pptx.util import Pt

            para = shape.text_frame.add_paragraph()
            run = para.add_run()
            run.text = paras[1]
            run.font.size = Pt(11)
            run.font.name = PPTX_FONT
            run.font.color.rgb = RGBColor(*color)

    minutes = f" | {s.minutes:g} min" if s.minutes is not None else ""
    header = f"Slide {s.number} of {rs.total} | {s.section}{minutes}"
    slide.notes_slide.notes_text_frame.text = f"{header}\n\n{s.notes}".strip()


def write_pptx(
    slides: list[RenderSlide], out: Path, fit: dict[int, dict[str, Any]] | None = None
) -> None:
    """A 16:9 PPTX: one slide per block, native text boxes, diagrams as pictures, speaker notes."""
    from pptx import Presentation

    prs = Presentation()
    prs.slide_width, prs.slide_height = _px(WIDTH), _px(HEIGHT)
    for rs in slides:
        _pptx_slide(prs, rs, (fit or {}).get(rs.slide.number))
    prs.core_properties.title = COURSE
    prs.save(str(out))


# -- command line -------------------------------------------------------------------------------


def _print_report(deck: Deck) -> int:
    totals: dict[str, float] = {}
    for s in deck.slides:
        totals[s.section] = totals.get(s.section, 0) + (s.minutes or 0)
    shown = ", ".join(f"{name} {value:g}" for name, value in totals.items())
    print(f"minutes per section: {shown}; total {sum(totals.values()):g}")
    pending, empty = open_items(deck)
    for number, where, p in pending:
        print(
            f"  slide {number:>2} {p.kind} ({where}): {p.what}"
            + (f" [{p.source}]" if p.source else "")
        )
    for number, name in empty:
        print(f"  slide {number:>2} empty result slot: {name}")
    marked = [(where, p) for _, where, p in pending if p.kind == "PENDING"]
    on_slides = sum(1 for where, _ in marked if where == "slide")
    captures = len(pending) - len(marked)
    print(
        f"open items: {len(marked)} PENDING markers ({on_slides} on slides, "
        f"{len(marked) - on_slides} in speaker notes), {captures} CAPTURE markers, "
        f"{len(empty)} empty result slots"
    )
    return len(pending) + len(empty)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--deck", type=Path, default=DECK, help="deck source (slides/deck.md)")
    parser.add_argument("--out", type=Path, default=OUT, help="output folder (slides/dist)")
    parser.add_argument("--expect", type=int, default=EXPECTED, help="required slide count")
    parser.add_argument("--check", action="store_true", help="validate only; no browser, no files")
    parser.add_argument("--strict", action="store_true", help="exit 1 if open items remain")
    args = parser.parse_args(argv)
    started = time.monotonic()
    deck_path = args.deck.resolve()
    try:
        deck = parse_deck(deck_path.read_text(encoding="utf-8"))
    except ParseError as exc:
        print(f"x {deck_path.name}: {exc}")
        return 2
    errors = validate(deck, deck_path.parent, args.expect)
    for error in errors:
        print(f"x {error}")
    if errors:
        print(f"{len(errors)} validation errors; nothing built")
        return 2
    print(f"{deck_path.name}: {len(deck.slides)} slides, valid")
    if not args.check:
        out = args.out.resolve()
        out.mkdir(parents=True, exist_ok=True)
        sources = [src for s in deck.slides for src in s.diagrams]
        try:
            pngs = render_diagrams(sources, out / "diagrams")
            slides = prepare(deck, deck_path.parent, REPO, pngs)
            html_path = out / "deck.html"
            html_path.write_text(render_html(slides), encoding="utf-8")
            fit = export_pdf(html_path, out / "deck.pdf")
        except BrowserUnavailable as exc:
            print(f"x {exc}")
            return 3
        except Exception as exc:  # noqa: BLE001 - mermaid or network errors end the build here
            print(f"x browser step failed: {type(exc).__name__}: {exc}")
            return 3
        write_pptx(slides, out / "deck.pptx", fit)
        for number, info in sorted(fit.items()):
            if info.get("overflow"):
                print(f"! slide {number}: text still overflows at {info['font_px']}px; shorten it")
            elif info.get("font_px", 28) < 22:
                print(f"! slide {number}: body text shrunk to {info['font_px']}px to fit")
        print(f"diagrams: {len(set(sources))} (cached under {out / 'diagrams'})")
        for name in ("deck.html", "deck.pdf", "deck.pptx"):
            size = (out / name).stat().st_size / 1024
            print(f"wrote {out / name} ({size:,.0f} KB)")
        print(f"pdf pages: {pdf_page_count(out / 'deck.pdf')}")
    remaining = _print_report(deck)
    print(f"build time: {time.monotonic() - started:.1f} s")
    if args.strict and remaining:
        print("x --strict: open items remain")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
